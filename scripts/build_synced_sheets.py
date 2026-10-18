#!/usr/bin/env python3
"""Build data/synced_sheets.csv from the archived perforated sheets.

Usage: python3 scripts/build_synced_sheets.py data/legacy/Perforated_Sheet.zip data/synced_sheets.csv

The archive holds the Word documents (VCI_Perforated_<P1>_<P2>.docx) used to print
specimen ID sheets before the ID Printer existed. Each folder is named
"<group> (<description>)". This script writes one row per prefix found in each
document, in the print log format, with source "sync":

  - print_id       S- followed by 8 hexadecimal characters derived from the file
                   path and prefix, so the same input always gives the same ID.
  - timestamp_utc  Last-modified time of the file in the archive (Google Drive
                   export, UTC). The actual print date is not recorded.
  - name, program  Taken from the folder description (see FOLDERS).
  - start, end     Lowest and highest number on the sheet. codes is the number of
                   distinct codes (the old sheets skip some numbers).
  - reprint        TRUE if an earlier row in this file already holds codes with
                   the same prefix (the same labels exist in more than one folder).
  - note           Archive path of the document, and any inconsistency found.

The Template folder is skipped. Requires python-docx (pip install python-docx).
"""

import csv
import hashlib
import io
import re
import sys
import zipfile
from datetime import datetime

import docx

HEADERS = [
    "print_id", "timestamp_utc", "name", "program", "prefix", "start", "end",
    "first_code", "last_code", "codes", "pages", "paper", "reprint", "source", "note",
]

# Folder description -> (name, program). Name is the printer where the folder states it.
FOLDERS = {
    "Printed, Mayuge Stationary": ("Mayuge Stationary", "Mayuge"),
    "Kenya, Printed at JHU": ("JHU", "Kenya"),
    "Colombia, PECET": ("Unknown", "Colombia (PECET)"),
    "VectorCam Exclusive Sheets": ("Unknown", "VectorCam"),
}
SKIP_FOLDERS = {"Template"}

CODE_RE = re.compile(r"\b([A-Z]{3})(\d{3})\b")
FILE_RE = re.compile(r"VCI_Perforated_([A-Z]{3})_([A-Z]{3})\.docx$")
FOLDER_RE = re.compile(r"^(\w+) \((.*?)\)+$")
GROUP_ORDER = {"Printed, Mayuge Stationary": 0}  # original sheets first, copies after


def paper_of(document) -> str:
    s = document.sections[0]
    w, h = round(s.page_width.mm), round(s.page_height.mm)
    if (w, h) == (216, 279):
        return "Letter"
    if (w, h) == (210, 297):
        return "A4"
    return f"{w}x{h}mm"


def read_entries(zip_path: str):
    entries = []
    with zipfile.ZipFile(zip_path) as z:
        for info in z.infolist():
            parts = info.filename.split("/")
            if not info.filename.endswith(".docx") or len(parts) < 3:
                continue
            folder, filename = parts[-2], parts[-1]
            if folder in SKIP_FOLDERS:
                continue
            m = FOLDER_RE.match(folder)
            if not m or m.group(2) not in FOLDERS:
                raise SystemExit(f"Unknown folder: {folder}")
            document = docx.Document(io.BytesIO(z.read(info)))
            numbers: dict[str, set[int]] = {}
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for prefix, num in CODE_RE.findall(cell.text):
                            numbers.setdefault(prefix, set()).add(int(num))
            pages = sum(1 for t in document.tables if any(CODE_RE.search(c.text) for c in t.rows[0].cells))
            entries.append({
                "path": info.filename,
                "folder": folder,
                "description": m.group(2),
                "filename": filename,
                "modified": datetime(*info.date_time).strftime("%Y-%m-%dT%H:%M:%SZ"),
                "numbers": numbers,
                "pages_per_prefix": max(1, pages // max(1, len(numbers))),
                "paper": paper_of(document),
            })
    # Original folders first, then copies, then by path, so duplicates are flagged on the copies.
    entries.sort(key=lambda e: (GROUP_ORDER.get(e["description"], 1), e["filename"].startswith("Copy of"), e["path"]))
    return entries


def build_rows(entries):
    rows, holders = [], {}
    for e in entries:
        name, program = FOLDERS[e["description"]]
        expected = FILE_RE.search(e["filename"])
        expected = set(expected.groups()) if expected else set()
        for prefix in sorted(e["numbers"]):
            nums = e["numbers"][prefix]
            notes = [e["path"]]
            if expected and prefix not in expected:
                notes.append(f"file name lists {'/'.join(sorted(expected))} but the sheet holds {prefix} codes")
            reprint = prefix in holders
            if reprint:
                notes.append("same codes also in " + "; ".join(holders[prefix]))
            holders.setdefault(prefix, []).append(e["path"])
            start, end = min(nums), max(nums)
            rows.append({
                "print_id": "S-" + hashlib.sha1(f"{e['path']}|{prefix}".encode()).hexdigest()[:8].upper(),
                "timestamp_utc": e["modified"],
                "name": name,
                "program": program,
                "prefix": prefix,
                "start": start,
                "end": end,
                "first_code": f"{prefix}{start:03d}",
                "last_code": f"{prefix}{end:03d}",
                "codes": len(nums),
                "pages": e["pages_per_prefix"],
                "paper": e["paper"],
                "reprint": "TRUE" if reprint else "FALSE",
                "source": "sync",
                "note": " | ".join(notes),
            })
    return rows


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__.strip().splitlines()[2], file=sys.stderr)
        return 2
    rows = build_rows(read_entries(sys.argv[1]))
    with open(sys.argv[2], "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HEADERS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    dups = sum(r["reprint"] == "TRUE" for r in rows)
    print(f"{sys.argv[2]}: {len(rows)} rows, {len({r['prefix'] for r in rows})} prefixes, {dups} duplicate rows.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
